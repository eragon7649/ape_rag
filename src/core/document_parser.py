# src/core/document_parser.py

import os
import asyncio
from typing import Dict, Any, List
from raganything import RAGAnything
from .config import RAG_CONFIG

class UnsupportedFormatError(Exception):
    """Exception khi format file không được hỗ trợ"""
    pass

class ParsingError(Exception):
    """Exception khi parsing thất bại"""
    pass

class EncodingError(Exception):
    """Exception khi không thể decode file"""
    pass

class DocumentParser:
    def __init__(self, rag_processor: RAGAnything):
        self.rag = rag_processor
        self.supported_formats = {
            '.pdf': self._parse_pdf,
            '.txt': self._parse_txt,
            '.docx': self._parse_docx,
            '.md': self._parse_markdown,
            '.html': self._parse_html
        }
    
    async def parse_document(self, file_path: str) -> Dict[str, Any]:
        """Parse document với error handling tốt hơn"""
        file_ext = os.path.splitext(file_path)[1].lower()
        file_base_name = os.path.splitext(os.path.basename(file_path))[0]
        
        print(f"🔍 Bắt đầu parse file: {file_path}")
        
        if file_ext not in self.supported_formats:
            raise UnsupportedFormatError(f"Format {file_ext} không được hỗ trợ. Các format hỗ trợ: {list(self.supported_formats.keys())}")
        
        try:
            # Thử parse với method chính
            result = await self.supported_formats[file_ext](file_path)
            print(f"✅ Parse thành công với method chính cho {file_base_name}")
            return result
        except Exception as e:
            print(f"⚠️ Method chính thất bại: {e}")
            # Fallback methods
            return await self._fallback_parse(file_path, file_ext, e)
    
    async def _parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """Parse PDF với multiple methods"""
        methods = ['auto', 'ocr', 'layout', 'table']
        file_base_name = os.path.splitext(os.path.basename(file_path))[0]
        
        for method in methods:
            try:
                print(f"🔄 Thử parse PDF với method: {method}")
                await self.rag.process_document_complete(
                    file_path=file_path,
                    output_dir=os.path.join(RAG_CONFIG.working_dir, "parsed_docs"),
                    parse_method=method
                )
                print(f"✅ Parse PDF thành công với method: {method}")
                return {'method': method, 'status': 'success'}
            except Exception as e:
                print(f"⚠️ Method {method} failed: {e}")
                continue
        
        raise ParsingError("Tất cả methods parse PDF đều thất bại")
    
    async def _parse_txt(self, file_path: str) -> Dict[str, Any]:
        """Parse TXT với encoding detection"""
        encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'ascii']
        file_base_name = os.path.splitext(os.path.basename(file_path))[0]
        
        for encoding in encodings:
            try:
                print(f"🔄 Thử đọc TXT với encoding: {encoding}")
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                print(f"✅ Đọc TXT thành công với encoding: {encoding}")
                return {
                    'content': content, 
                    'encoding': encoding,
                    'method': 'direct_read',
                    'status': 'success'
                }
            except UnicodeDecodeError:
                print(f"⚠️ Encoding {encoding} failed")
                continue
            except Exception as e:
                print(f"⚠️ Lỗi đọc file với {encoding}: {e}")
                continue
        
        raise EncodingError("Không thể decode file với các encoding đã thử")
    
    async def _parse_docx(self, file_path: str) -> Dict[str, Any]:
        """Parse DOCX file"""
        try:
            from docx import Document
            doc = Document(file_path)
            content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            # Lấy nội dung từ tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        content.append(" | ".join(row_text))
            
            return {
                'content': '\n'.join(content),
                'method': 'python_docx',
                'status': 'success'
            }
        except ImportError:
            raise UnsupportedFormatError("python-docx không được cài đặt. Chạy: pip install python-docx")
        except Exception as e:
            raise ParsingError(f"Lỗi parse DOCX: {e}")
    
    async def _parse_markdown(self, file_path: str) -> Dict[str, Any]:
        """Parse Markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            return {
                'content': content,
                'method': 'direct_read',
                'status': 'success'
            }
        except Exception as e:
            raise ParsingError(f"Lỗi parse Markdown: {e}")
    
    async def _parse_html(self, file_path: str) -> Dict[str, Any]:
        """Parse HTML file"""
        try:
            from bs4 import BeautifulSoup
            with open(file_path, 'r', encoding='utf-8') as f:
                html_content = f.read()
            
            soup = BeautifulSoup(html_content, 'html.parser')
            # Loại bỏ script và style
            for script in soup(["script", "style"]):
                script.decompose()
            
            text = soup.get_text()
            # Làm sạch text
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return {
                'content': text,
                'method': 'beautifulsoup',
                'status': 'success'
            }
        except ImportError:
            raise UnsupportedFormatError("beautifulsoup4 không được cài đặt. Chạy: pip install beautifulsoup4")
        except Exception as e:
            raise ParsingError(f"Lỗi parse HTML: {e}")
    
    async def _fallback_parse(self, file_path: str, file_ext: str, original_error: Exception) -> Dict[str, Any]:
        """Fallback parsing methods"""
        file_base_name = os.path.splitext(os.path.basename(file_path))[0]
        print(f"🔄 Thử fallback methods cho {file_base_name}")
        
        # Fallback 1: Thử đọc như text file
        if file_ext != '.txt':
            try:
                print("🔄 Fallback: Thử đọc như text file")
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
                return {
                    'content': content,
                    'method': 'fallback_text',
                    'status': 'partial_success',
                    'original_error': str(original_error)
                }
            except Exception as e:
                print(f"⚠️ Fallback text failed: {e}")
        
        # Fallback 2: Thử với binary mode
        try:
            print("🔄 Fallback: Thử đọc binary mode")
            with open(file_path, 'rb') as f:
                content = f.read()
            # Chỉ lấy phần text có thể đọc được
            text_content = ''.join(chr(b) for b in content if 32 <= b <= 126)
            return {
                'content': text_content,
                'method': 'fallback_binary',
                'status': 'partial_success',
                'original_error': str(original_error)
            }
        except Exception as e:
            print(f"⚠️ Fallback binary failed: {e}")
        
        # Nếu tất cả đều thất bại
        raise ParsingError(f"Tất cả methods parse đều thất bại. Original error: {original_error}")
    
    def get_supported_formats(self) -> List[str]:
        """Lấy danh sách format được hỗ trợ"""
        return list(self.supported_formats.keys())
    
    def is_supported_format(self, file_path: str) -> bool:
        """Kiểm tra format có được hỗ trợ không"""
        file_ext = os.path.splitext(file_path)[1].lower()
        return file_ext in self.supported_formats
